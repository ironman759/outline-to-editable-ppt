#!/usr/bin/env python3
"""Check package integrity and known hazards; real PowerPoint QA remains required."""
import argparse
import json
import posixpath
import zipfile
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from pptx import Presentation
from sync_scripts import check_mapping, check_scripts, load_manifest, walk_shapes

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def inspect(path, allow_animation=False):
    errors = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if len(names) != len(set(names)):
            errors.append("Duplicate ZIP entries")
        bad = z.testzip()
        if bad:
            errors.append(f"Bad ZIP CRC: {bad}")
        entries = set(names)
        for name in names:
            if name.endswith((".xml", ".rels")):
                root = ET.fromstring(z.read(name))
                for para in root.findall(".//a:p", NS):
                    if len(para.findall("a:pPr", NS)) > 1:
                        errors.append(f"Duplicate paragraph properties: {name}")
                if name.endswith(".rels"):
                    base = posixpath.dirname(posixpath.dirname(name))
                    ids = []
                    for rel in root:
                        ids.append(rel.attrib["Id"])
                        if rel.get("TargetMode") == "External":
                            continue
                        target = unquote(rel.attrib["Target"])
                        dest = posixpath.normpath(posixpath.join(base, target)).lstrip("/") if not target.startswith("/") else target.lstrip("/")
                        if dest not in entries:
                            errors.append(f"Missing relationship target: {name} → {dest}")
                    if len(ids) != len(set(ids)):
                        errors.append(f"Duplicate relationship ID: {name}")
                if name == "ppt/presentation.xml":
                    order = [x.tag.split("}")[-1] for x in root]
                    if "notesMasterIdLst" in order and "sldIdLst" in order and order.index("notesMasterIdLst") > order.index("sldIdLst"):
                        errors.append("notesMasterIdLst must precede sldIdLst")
                if name.startswith("ppt/slides/slide") and not allow_animation:
                    if root.find("p:transition", NS) is not None or root.find("p:timing", NS) is not None:
                        errors.append(f"Unexpected animation or transition: {name}")
    prs = Presentation(path)
    counts = []
    tolerance = 9144  # 0.01 inch: tolerate rounding only
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(walk_shapes(slide.shapes))
        native = sum((s.has_text_frame and bool(s.text.strip())) or s.has_table for s in shapes)
        if not native:
            errors.append(f"Slide {i} has no editable text/table; inspect flattening")
        for s in slide.shapes:
            if min(s.left, s.top) < -tolerance or s.left + s.width > prs.slide_width + tolerance or s.top + s.height > prs.slide_height + tolerance:
                errors.append(f"Out-of-bounds shape on slide {i}: {s.name}")
        counts.append({"slide": i, "shapes": len(shapes), "nonempty_text_shapes": sum(s.has_text_frame and bool(s.text.strip()) for s in shapes), "notes": slide.has_notes_slide})
    return errors, counts


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx")
    ap.add_argument("--manifest")
    ap.add_argument("--docx")
    ap.add_argument("--allow-animation", action="store_true")
    args = ap.parse_args()
    if args.docx and not args.manifest:
        ap.error("--docx requires --manifest")
    errors, counts = inspect(args.pptx, args.allow_animation)
    if args.manifest:
        data = load_manifest(args.manifest)
        check_mapping(Presentation(args.pptx), data)
        for i, (slide, page) in enumerate(zip(Presentation(args.pptx).slides, data["slides"]), 1):
            if not slide.has_notes_slide or slide.notes_slide.notes_text_frame.text != page["script"]:
                errors.append(f"Notes mismatch: slide {i}")
        if args.docx:
            check_scripts(args.pptx, args.docx, data)
    print(json.dumps({"package_checks": "FAIL" if errors else "PASS", "errors": errors, "slides": counts,
                      "powerpoint_open_test": "REQUIRED_SEPARATELY"}, ensure_ascii=False, indent=2))
    raise SystemExit(bool(errors))


if __name__ == "__main__":
    main()
