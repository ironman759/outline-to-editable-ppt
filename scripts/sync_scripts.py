#!/usr/bin/env python3
"""Synchronize saved native slides and a page-mapped Word from one script source."""
import argparse
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt as WordPt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def load_manifest(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    pages = data["slides"]
    if not pages:
        raise ValueError("manifest slides is empty")
    ids = []
    for p in pages:
        for key in ("id", "title", "script"):
            if not isinstance(p.get(key), str) or not p[key].strip():
                raise ValueError(f"Missing/non-text {key}: {p.get('id')}")
        if not isinstance(p.get("seconds", 0), (int, float)) or p.get("seconds", 0) < 0:
            raise ValueError(f"Invalid seconds: {p['id']}")
        ids.append(p["id"])
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate stable slide IDs")
    return data


def check_mapping(prs, data):
    if len(prs.slides) != len(data["slides"]):
        raise ValueError("PPT and manifest slide counts differ")
    for number, (slide, page) in enumerate(zip(prs.slides, data["slides"]), 1):
        if slide._element.cSld.get("name") != f"page-id:{page['id']}":
            raise ValueError(f"Slide {number} stable ID does not match {page['id']}; verify page mapping before stamping IDs")
        texts = [s.text for s in walk_shapes(slide.shapes) if s.has_text_frame]
        if page["title"] not in texts:
            raise ValueError(f"Slide {number} title does not match {page['id']}")


def check_scripts(pptx_path, docx_path, data):
    prs = Presentation(pptx_path)
    check_mapping(prs, data)
    paras = [p.text for p in Document(docx_path).paragraphs]
    body = []
    for i, (slide, page) in enumerate(zip(prs.slides, data["slides"]), 1):
        if not slide.has_notes_slide or slide.notes_slide.notes_text_frame.text != page["script"]:
            raise ValueError(f"Notes mismatch on slide {i}")
        heading = f"第 {i:02d} 页｜{page['title']}"
        body.extend([heading, f"ID：{page['id']}｜部分：{page.get('section', '')}｜预计 {page.get('seconds', 0):g} 秒", page["script"]])
    if paras[2:] != body:
        raise ValueError("Word page order, IDs, titles, timing or scripts differ from manifest")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("source")
    ap.add_argument("manifest")
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--docx", required=True)
    ap.add_argument("--force", action="store_true", help="Replace previously archived outputs")
    args = ap.parse_args()
    source = Path(args.source).resolve()
    outputs = [Path(args.pptx).resolve(), Path(args.docx).resolve()]
    if len(set(outputs + [source, Path(args.manifest).resolve()])) != 4:
        raise ValueError("Source, manifest and outputs must be separate files")
    for out in outputs:
        if out.exists() and not args.force:
            raise FileExistsError(f"Archive old version first; output exists: {out}")
    data = load_manifest(args.manifest)
    prs = Presentation(source)
    check_mapping(prs, data)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = data.get("font", "Microsoft YaHei")
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style.font.name)
    doc.styles["Normal"].font.size = WordPt(11)
    doc.styles["Normal"].paragraph_format.space_after = WordPt(7)
    doc.styles["Heading 1"].font.size = WordPt(16)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.add_paragraph(data.get("title", "逐页演讲稿"), "Title")
    total = sum(p.get("seconds", 0) for p in data["slides"])
    doc.add_paragraph(f"共 {len(prs.slides)} 页；预计 {total / 60:g} 分钟（排练预算，非实测时长）。与 PPT 备注使用同一讲稿源。")
    for i, (slide, page) in enumerate(zip(prs.slides, data["slides"]), 1):
        slide.notes_slide.notes_text_frame.text = page["script"]
        h = doc.add_paragraph(f"第 {i:02d} 页｜{page['title']}", "Heading 1")
        if i > 1:
            h.paragraph_format.page_break_before = True
        doc.add_paragraph(f"ID：{page['id']}｜部分：{page.get('section', '')}｜预计 {page.get('seconds', 0):g} 秒")
        doc.add_paragraph(page["script"])
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for out in outputs:
        out.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=outputs[0].parent) as tmp:
        test_pptx, test_docx = Path(tmp) / "synced.pptx", Path(tmp) / "script.docx"
        prs.save(test_pptx)
        doc.save(test_docx)
        check_scripts(test_pptx, test_docx, data)
        # Stage both files before committing; callers must check both after a write error.
        import shutil
        with TemporaryDirectory(dir=outputs[1].parent) as word_tmp:
            word_file = Path(word_tmp) / "script.docx"
            shutil.copyfile(test_docx, word_file)
            os.replace(test_pptx, outputs[0])
            os.replace(word_file, outputs[1])
    check_scripts(outputs[0], outputs[1], data)
    print(json.dumps({"slides": len(prs.slides), "scripts_readback": "PASS", "pptx": str(outputs[0]), "docx": str(outputs[1])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
